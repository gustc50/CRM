"""Plano de contas e conciliação contábil automática.

O contador importa o plano de contas da empresa, exportado do sistema
contábil dele (Domínio, Questor, Alterdata, Calima...), e liga as coisas do
dia a dia às contas: cada categoria a uma conta de resultado, cada banco a
uma conta do disponível, palavras da descrição a contas específicas. Com
isso, todo lançamento ganha sozinho o seu débito e o seu crédito:

    conta a pagar:   D despesa (categoria/regra)   C banco (ou Caixa)
    conta a receber: D banco (ou Caixa)            C receita (categoria/regra)

A classificação é calculada na hora, a partir dos vínculos: mudou um
vínculo, todos os lançamentos já aparecem com a conta nova — não há nada
gravado para ficar velho.

Sobre os arquivos: cada sistema exporta o plano num leiaute um pouco
diferente (colunas em outra ordem, com ou sem cabeçalho, classificação com
pontos ou só com dígitos e máscara). O leitor não depende de posição fixa:
acha as colunas pelo nome do cabeçalho e, sem cabeçalho, pelo conteúdo — a
coluna com 1.1.1.01.001 é a classificação, a de números curtos é o código
reduzido, a de texto é a descrição, a de S/A é o tipo.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

from models import (
    MAP_BANCO, MAP_CAIXA, MAP_CATEGORIA, MAP_DESPESA, MAP_RECEITA, MAP_TEXTO,
    ContaContabil, MapeamentoContabil, db,
)


class ErroPlano(Exception):
    pass


def normalizar(texto):
    """Sem acento, minúsculo e com um espaço só entre palavras."""
    sem_acento = unicodedata.normalize('NFKD', str(texto or ''))
    sem_acento = ''.join(c for c in sem_acento if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', sem_acento).strip().lower()


# ------------------------------------------------------------------ #
# Leitura do arquivo
# ------------------------------------------------------------------ #

# Nomes de coluna que cada sistema costuma usar (comparados sem acento)
SINONIMOS = {
    'classificacao': ('classificacao', 'classif', 'class.', 'estrutura', 'mascara', 'conta contabil',
                      'codigo contabil', 'cod. contabil', 'numero da conta', 'nr. conta', 'nº conta'),
    'codigo': ('reduzido', 'cod. reduzido', 'codigo reduzido', 'cod reduzido', 'cod.red', 'red.',
               'codigo', 'cod.', 'cod', 'conta'),
    'descricao': ('descricao', 'nome', 'titulo', 'denominacao', 'nome da conta', 'descricao da conta',
                  'historico da conta'),
    'tipo': ('tipo', 'a/s', 's/a', 'analitica/sintetica', 'sintetica/analitica', 'tipo da conta'),
    'natureza': ('natureza', 'nat.', 'nat', 'd/c', 'natureza da conta'),
}

_RE_CLASSIFICACAO = re.compile(r'^\d+(?:[.\-]\d+)+$')
_RE_INTEIRO = re.compile(r'^\d{1,12}$')
_TIPOS = {'s': 'S', 'a': 'A', 't': 'S', 'sintetica': 'S', 'analitica': 'A',
          'sintetico': 'S', 'analitico': 'A', 'titulo': 'S'}
_NATUREZAS = {'d': 'D', 'c': 'C', 'devedora': 'D', 'credora': 'C', 'dev': 'D', 'cred': 'C',
              'devedor': 'D', 'credor': 'C'}


def _decodificar(conteudo):
    for codificacao in ('utf-8-sig', 'cp1252'):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode('latin-1')


def _celula(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _linhas_do_arquivo(conteudo, nome_arquivo):
    """O arquivo vira uma lista de linhas, cada uma uma lista de células."""
    if conteudo[:2] == b'PK' or nome_arquivo.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
        except Exception as exc:
            raise ErroPlano(f'Não deu para abrir a planilha: {exc}') from exc
        return [[_celula(c) for c in linha] for linha in planilha.iter_rows(values_only=True)]
    if nome_arquivo.lower().endswith('.xls'):
        raise ErroPlano('Arquivo .xls (Excel antigo) não é lido: salve como .xlsx ou .csv e envie de novo.')

    texto = _decodificar(conteudo)
    linhas = [l for l in texto.splitlines() if l.strip()]
    if not linhas:
        raise ErroPlano('O arquivo está vazio.')

    # Separador: o que aparece o mesmo número de vezes na maioria das linhas
    amostra = linhas[:80]
    for separador in (';', '\t', '|', ','):
        contagens = Counter(l.count(separador) for l in amostra)
        vezes, quantas = contagens.most_common(1)[0]
        if vezes >= 1 and quantas / len(amostra) >= 0.6:
            return [[c.strip() for c in linha] for linha in csv.reader(linhas, delimiter=separador)]

    # Relatório em texto, colunas alinhadas por espaço
    return [[c.strip() for c in re.split(r'\s{2,}|\t', l.strip())] for l in linhas]


def _campo_do_cabecalho(texto):
    alvo = normalizar(texto).strip(' :')
    if not alvo:
        return None
    for campo, nomes in SINONIMOS.items():
        if alvo in nomes:
            return campo
    for campo, nomes in SINONIMOS.items():
        if any(len(n) > 3 and n in alvo for n in nomes):
            return campo
    return None


def _achar_cabecalho(linhas):
    for i, linha in enumerate(linhas[:15]):
        campos = {}
        for j, celula in enumerate(linha):
            campo = _campo_do_cabecalho(celula)
            if campo and campo not in campos.values():
                campos[j] = campo
        if len(campos) >= 2 and ({'descricao', 'classificacao'} & set(campos.values())):
            return i, campos
    return None, {}


def _perfil_das_colunas(dados):
    """Como é o conteúdo de cada coluna: serve para achar quem é quem sem cabeçalho."""
    largura = max((len(l) for l in dados), default=0)
    perfil = []
    for j in range(largura):
        valores = [l[j] for l in dados if j < len(l) and l[j]]
        n = len(valores) or 1
        perfil.append({
            'classificacao': sum(bool(_RE_CLASSIFICACAO.match(v)) for v in valores) / n,
            'inteiro': sum(bool(_RE_INTEIRO.match(v)) for v in valores) / n,
            'tipo': sum(normalizar(v) in _TIPOS for v in valores) / n,
            'natureza': sum(normalizar(v) in _NATUREZAS for v in valores) / n,
            'letras': sum(sum(ch.isalpha() for ch in v) for v in valores) / n,
            'max_digitos': max((len(v) for v in valores if _RE_INTEIRO.match(v)), default=0),
            'preenchidas': len(valores),
        })
    return perfil


def _segmentos_da_mascara(mascara):
    partes = [p for p in re.split(r'[.\-]', (mascara or '').strip()) if p]
    return [len(p) for p in partes] if partes and all(re.fullmatch(r'[9#0Xx]+', p) for p in partes) else []


def _aplicar_mascara(digitos, segmentos):
    """11101001 com 9.9.9.99.999 vira 1.1.1.01.001; zeros no fim são de conta-mãe."""
    partes, pos = [], 0
    for tamanho in segmentos:
        if pos >= len(digitos):
            break
        partes.append(digitos[pos:pos + tamanho])
        pos += tamanho
    if pos < len(digitos):
        partes.append(digitos[pos:])
    while len(partes) > 1 and set(partes[-1]) == {'0'}:
        partes.pop()
    return '.'.join(partes)


def ler_plano(conteudo, nome_arquivo='', mascara=''):
    """Lê o plano de contas. Devolve (contas, resumo).

    `contas` é uma lista de dicionários com classificacao, codigo, descricao,
    tipo (A/S), natureza (D/C ou None) e nivel.
    """
    linhas = _linhas_do_arquivo(conteudo, nome_arquivo)
    i_cab, campos = _achar_cabecalho(linhas)
    dados = linhas[i_cab + 1:] if i_cab is not None else linhas
    perfil = _perfil_das_colunas(dados)
    if not perfil:
        raise ErroPlano('Não achei nenhuma coluna no arquivo.')

    coluna = {campo: j for j, campo in campos.items()}

    def melhor(chave, minimo, excluir=()):
        candidatas = [(p[chave], j) for j, p in enumerate(perfil)
                      if j not in excluir and p[chave] >= minimo]
        return max(candidatas)[1] if candidatas else None

    # Classificação: a coluna com 1.1.1.01.001. Se o cabeçalho apontou outra
    # (ou não havia cabeçalho), o conteúdo decide.
    j_class = melhor('classificacao', 0.5)
    if j_class is not None:
        if coluna.get('codigo') == j_class:  # "Código" que na verdade é a classificação
            coluna.pop('codigo')
        coluna['classificacao'] = j_class
    elif 'classificacao' not in coluna:
        # Sem pontos: a coluna de números mais longos é a classificação
        inteiros = sorted(((p['max_digitos'], j) for j, p in enumerate(perfil) if p['inteiro'] >= 0.6),
                          reverse=True)
        if inteiros:
            coluna['classificacao'] = inteiros[0][1]
            if len(inteiros) > 1 and 'codigo' not in coluna:
                coluna['codigo'] = inteiros[1][1]

    usadas = set(coluna.values())
    if 'codigo' not in coluna:
        j = melhor('inteiro', 0.6, excluir=usadas)
        if j is not None:
            coluna['codigo'] = j
            usadas.add(j)
    if 'descricao' not in coluna:
        j = melhor('letras', 2, excluir=usadas)
        if j is not None:
            coluna['descricao'] = j
            usadas.add(j)
    if 'tipo' not in coluna:
        j = melhor('tipo', 0.8, excluir=usadas)
        if j is not None:
            coluna['tipo'] = j
            usadas.add(j)
    if 'natureza' not in coluna:
        j = melhor('natureza', 0.8, excluir=usadas)
        if j is not None:
            coluna['natureza'] = j

    if 'classificacao' not in coluna or 'descricao' not in coluna:
        raise ErroPlano('Não achei as colunas de classificação e de descrição das contas. Confira se '
                        'o arquivo é o plano de contas exportado do sistema contábil.')

    segmentos = _segmentos_da_mascara(mascara)

    def valor(linha, campo):
        j = coluna.get(campo)
        return linha[j].strip() if j is not None and j < len(linha) else ''

    contas, ignoradas = {}, 0
    for linha in dados:
        classificacao = valor(linha, 'classificacao').replace(' ', '').replace('-', '.').strip('.')
        descricao = re.sub(r'\s+', ' ', valor(linha, 'descricao'))
        if not re.fullmatch(r'\d+(?:\.\d+)*', classificacao) or not re.search(r'[A-Za-zÀ-ÿ]', descricao):
            ignoradas += 1  # cabeçalho de página, total, linha em branco
            continue
        if '.' not in classificacao and segmentos:
            classificacao = _aplicar_mascara(classificacao, segmentos)
        codigo = valor(linha, 'codigo')
        contas[classificacao] = {
            'classificacao': classificacao[:40],
            'codigo': (codigo if _RE_INTEIRO.match(codigo) else '')[:20] or None,
            'descricao': descricao[:150],
            'tipo': _TIPOS.get(normalizar(valor(linha, 'tipo'))),
            'natureza': _NATUREZAS.get(normalizar(valor(linha, 'natureza'))),
        }

    if not contas:
        raise ErroPlano('Nenhuma conta foi reconhecida no arquivo.')

    # Tipo que o arquivo não trouxe: quem tem conta "filha" é sintética
    com_filhas = set()
    for classificacao in contas:
        if '.' in classificacao:
            partes = classificacao.split('.')
            com_filhas.update('.'.join(partes[:i]) for i in range(1, len(partes)))
        else:
            com_filhas.update(classificacao[:i] for i in range(1, len(classificacao)))
    for c in contas.values():
        if c['tipo'] is None:
            c['tipo'] = 'S' if c['classificacao'] in com_filhas else 'A'
        c['nivel'] = c['classificacao'].count('.') + 1 if '.' in c['classificacao'] else len(c['classificacao'])

    lista = sorted(contas.values(), key=lambda c: [int(p) for p in c['classificacao'].split('.')])
    resumo = {
        'contas': len(lista),
        'analiticas': sum(c['tipo'] == 'A' for c in lista),
        'ignoradas': ignoradas,
        'colunas': {campo: j + 1 for campo, j in coluna.items()},
        'com_cabecalho': i_cab is not None,
    }
    return lista, resumo


def gravar_plano(empresa_id, contas):
    """Substitui o plano da empresa, mantendo os vínculos de quem continua.

    A conta é reconhecida pela classificação: reimportar o mesmo plano
    atualizado não desfaz o de-para. Conta que sumiu do arquivo é apagada —
    a não ser que esteja num vínculo; aí fica inativa e a tela avisa.
    """
    existentes = {c.classificacao: c for c in ContaContabil.query.filter_by(empresa_id=empresa_id)}
    em_uso = {m.conta_id for m in MapeamentoContabil.query.filter_by(empresa_id=empresa_id)}
    novas = atualizadas = removidas = inativas = 0
    for dados in contas:
        conta = existentes.pop(dados['classificacao'], None)
        if conta is None:
            db.session.add(ContaContabil(empresa_id=empresa_id, **dados))
            novas += 1
        else:
            for campo, valor_ in dados.items():
                setattr(conta, campo, valor_)
            conta.ativa = True
            atualizadas += 1
    for conta in existentes.values():
        if conta.id in em_uso:
            conta.ativa = False
            inativas += 1
        else:
            db.session.delete(conta)
            removidas += 1
    db.session.commit()
    return {'novas': novas, 'atualizadas': atualizadas, 'removidas': removidas, 'inativas': inativas}


# ------------------------------------------------------------------ #
# Sugestão automática de vínculos (logo depois da importação)
# ------------------------------------------------------------------ #

_SEM_PESO = {'de', 'da', 'do', 'das', 'dos', 'e', 'a', 'o', 'as', 'os', 'com', 'para', 'por', 'em',
             'sa', 'ltda', 'me', 'banco', 'conta', 'contas', 'outras', 'outros', 'diversas', 'diversos',
             'geral', 'gerais'}
# Contas que casam pelo nome mas quase nunca são o que a categoria quer dizer
_REDUTORAS = {'devolucao', 'devolucoes', 'desconto', 'descontos', 'cancelamento', 'cancelamentos',
              'abatimento', 'abatimentos', 'deducao', 'deducoes', 'provisao', 'provisoes', 'a pagar',
              'a receber', 'acumulada', 'depreciacao'}


def _singular(palavra):
    for fim, troca in (('oes', 'ao'), ('aes', 'ao'), ('eis', 'el'), ('ais', 'al'), ('is', 'l')):
        if palavra.endswith(fim) and len(palavra) > len(fim) + 2:
            return palavra[:-len(fim)] + troca
    if palavra.endswith('s') and len(palavra) > 3:
        return palavra[:-1]
    return palavra


def palavras(texto):
    return {_singular(p) for p in re.findall(r'[a-z0-9]+', normalizar(texto)) if p not in _SEM_PESO}


def _melhor_conta(nome, contas, tipo_lancamento=None):
    alvo = palavras(nome)
    if not alvo:
        return None
    melhor, nota_melhor = None, (0, 0, 0)
    for conta in contas:
        tem = palavras(conta.descricao)
        comum = alvo & tem
        if not comum:
            continue
        cobertura = len(comum) / len(alvo)
        if cobertura < 0.6:
            continue
        descricao = normalizar(conta.descricao)
        penalidade = 1 if any(r in descricao for r in _REDUTORAS) else 0
        natureza_ok = 1 if (conta.natureza is None
                            or (tipo_lancamento == 'Pagar' and conta.natureza == 'D')
                            or (tipo_lancamento == 'Receber' and conta.natureza == 'C')) else 0
        nota = (cobertura - penalidade, natureza_ok, len(comum) / len(tem))
        if nota > nota_melhor:
            melhor, nota_melhor = conta, nota
    return melhor if nota_melhor[0] >= 0.6 else None


def sugerir_vinculos(empresa_id, categorias, bancos):
    """Liga pelo nome o que ainda não tem vínculo, para a conciliação já
    começar a funcionar logo depois da importação. O contador revisa: tudo o
    que vem daqui aparece como "sugerido".

    `bancos` é uma lista de (id, nome da conta, nome do banco).
    """
    contas = ContaContabil.query.filter_by(empresa_id=empresa_id, tipo='A', ativa=True).all()
    ja = {(m.tipo, m.referencia_id) for m in MapeamentoContabil.query.filter_by(empresa_id=empresa_id)}
    criados = 0

    def criar(tipo, referencia_id, conta):
        nonlocal criados
        db.session.add(MapeamentoContabil(empresa_id=empresa_id, tipo=tipo, referencia_id=referencia_id,
                                          conta_id=conta.id, origem='sugestao'))
        criados += 1

    for categoria in categorias:
        if (MAP_CATEGORIA, categoria.id) in ja:
            continue
        conta = _melhor_conta(categoria.nome, contas, categoria.tipo)
        if conta is not None:
            criar(MAP_CATEGORIA, categoria.id, conta)

    for banco_id, nome_conta, nome_banco in bancos:
        if (MAP_BANCO, banco_id) in ja:
            continue
        conta = (_melhor_conta(nome_banco, contas) if nome_banco else None) or _melhor_conta(nome_conta, contas)
        if conta is not None:
            criar(MAP_BANCO, banco_id, conta)

    if (MAP_CAIXA, None) not in ja:
        caixas = [c for c in contas if re.fullmatch(r'caixa( geral)?', normalizar(c.descricao))]
        if caixas:
            criar(MAP_CAIXA, None, caixas[0])

    db.session.commit()
    return criados


# ------------------------------------------------------------------ #
# Conciliação: débito e crédito de cada lançamento
# ------------------------------------------------------------------ #

@dataclass
class Classificacao:
    debito: ContaContabil | None = None
    credito: ContaContabil | None = None
    origem: str = ''            # o que decidiu a conta de resultado
    pendencias: list = field(default_factory=list)

    @property
    def conciliado(self):
        return self.debito is not None and self.credito is not None


class Conciliador:
    """Carrega o de-para da empresa uma vez e classifica quantos lançamentos precisar."""

    def __init__(self, empresa_id):
        self.contas = {c.id: c for c in ContaContabil.query.filter_by(empresa_id=empresa_id)}
        self.por_categoria, self.por_banco, self.textos, self.padroes = {}, {}, [], {}
        for m in (MapeamentoContabil.query.filter_by(empresa_id=empresa_id)
                  .order_by(MapeamentoContabil.id)):
            conta = self.contas.get(m.conta_id)
            if conta is None:
                continue
            if m.tipo == MAP_CATEGORIA:
                self.por_categoria[m.referencia_id] = conta
            elif m.tipo == MAP_BANCO:
                self.por_banco[m.referencia_id] = conta
            elif m.tipo == MAP_TEXTO and m.texto:
                self.textos.append((normalizar(m.texto), m.tipo_lancamento or '', conta, m.texto))
            else:
                self.padroes[m.tipo] = conta

    @property
    def tem_plano(self):
        return bool(self.contas)

    def _resultado(self, t):
        nome = ''
        if t.tipo == 'Pagar' and t.fornecedor is not None:
            nome = t.fornecedor.nome
        elif t.tipo == 'Receber' and t.cliente is not None:
            nome = t.cliente.nome
        alvo = normalizar(f'{t.descricao} {nome}')
        for texto, tipo, conta, original in self.textos:
            if (not tipo or tipo == t.tipo) and texto in alvo:
                return conta, f'regra "{original}"'
        if t.categoria_id in self.por_categoria:
            return self.por_categoria[t.categoria_id], 'categoria'
        padrao = self.padroes.get(MAP_DESPESA if t.tipo == 'Pagar' else MAP_RECEITA)
        if padrao is not None:
            return padrao, 'conta padrão'
        return None, ''

    def classificar(self, t):
        r = Classificacao()
        resultado, r.origem = self._resultado(t)
        if resultado is None:
            if t.categoria is not None:
                r.pendencias.append(f'categoria "{t.categoria.nome}" sem conta')
            else:
                r.pendencias.append('sem categoria nem regra')

        if t.conta_bancaria_id:
            financeira = self.por_banco.get(t.conta_bancaria_id)
            if financeira is None:
                r.pendencias.append(f'banco "{t.conta_bancaria.nome}" sem conta')
        else:
            financeira = self.padroes.get(MAP_CAIXA)
            if financeira is None:
                r.pendencias.append('sem banco e sem conta Caixa')

        if t.tipo == 'Pagar':
            r.debito, r.credito = resultado, financeira
        else:
            r.debito, r.credito = financeira, resultado
        return r

    def resumo(self, transacoes):
        """Quantos estão conciliados e, dos que não estão, por quê (agrupado)."""
        conciliados, motivos = 0, Counter()
        for t in transacoes:
            r = self.classificar(t)
            if r.conciliado:
                conciliados += 1
            motivos.update(r.pendencias)
        return {'total': len(transacoes), 'conciliados': conciliados,
                'pendentes': len(transacoes) - conciliados, 'motivos': motivos.most_common()}
